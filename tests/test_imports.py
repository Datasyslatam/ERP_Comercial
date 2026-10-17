"""Importación: plantillas, validación previa, duplicados, corrección y reintento, saldos iniciales aprobados."""
import io
import json

from helpers import TODAY, D, entries_for, scalar
from openpyxl import load_workbook


def _upload(api, entity, content: str, filename="datos.csv", **form):
    data = {"entity": entity, **{k: str(v) for k, v in form.items()}}
    return api.client.post("/api/v1/imports/jobs", headers=api.headers, data=data,
                           files={"file": (filename, content.encode("utf-8"), "text/csv")})


def test_template_download(env):
    r = env.client.get("/api/v1/imports/templates/items", headers=env.api("admin1").headers)
    wb = load_workbook(io.BytesIO(r.content))
    assert [c.value for c in wb["datos"][1]][:4] == ["sku", "name", "item_type", "base_uom_code"]
    assert wb["instrucciones"].max_row > 5


def test_parties_import_with_errors_fix_and_retry(env):
    a = env.api("admin1")
    csv_data = ("codigo;razon;tipo;nit;dv;cliente;proveedor\n"
                "C100;Cliente Uno SAS;NIT;900555444;{dv};SI;NO\n"
                "C101;Cliente Dos;NIT;900555445;9;SI;NO\n"            # DV incorrecto
                "C100;Duplicado en archivo;CC;12345678;;SI;NO\n"      # código repetido
                "C103;Sin rol;CC;87654321;;NO;NO\n").format(dv=__import__(
                    "app.modules.masters.validation", fromlist=["x"]).nit_check_digit("900555444"))
    mapping = {"codigo": "code", "razon": "legal_name", "tipo": "id_type", "nit": "tax_id", "dv": "check_digit",
               "cliente": "is_customer", "proveedor": "is_supplier"}
    r = _upload(a, "parties", csv_data, mapping=json.dumps(mapping))
    assert r.status_code == 201, r.text
    job = r.json()
    assert job["total_rows"] == 4 and job["valid_rows"] == 1 and job["invalid_rows"] == 3
    rows = a.get(f"/imports/jobs/{job['id']}/rows?status=invalid")
    errors = {x["row_number"]: " ".join(x["errors"]) for x in rows}
    assert "verificación" in errors[3] and "Duplicado" in errors[4] and "cliente, proveedor" in errors[5]
    # Corrección de filas y reintento
    from app.modules.masters.validation import nit_check_digit

    a.patch(f"/imports/rows/{rows[0]['id']}", {"data": {"check_digit": nit_check_digit("900555445")}})
    a.patch(f"/imports/rows/{rows[1]['id']}", {"data": {"code": "C102"}})
    job = a.post(f"/imports/jobs/{job['id']}/commit")
    assert job["imported_rows"] == 3 and job["status"] == "partially_imported"
    a.patch(f"/imports/rows/{rows[2]['id']}", {"data": {"is_customer": "SI"}})
    job = a.post(f"/imports/jobs/{job['id']}/commit")
    assert job["imported_rows"] == 4 and job["status"] == "imported"
    rec = a.get(f"/imports/jobs/{job['id']}/reconciliation")
    assert rec["ok"] and rec["persisted"] == 4
    # Volver a importar el mismo archivo detecta duplicados contra la base
    r2 = _upload(a, "parties", csv_data, mapping=json.dumps(mapping))
    assert r2.json()["valid_rows"] == 0


def test_opening_stock_requires_documented_approved_batch(env):
    a, b = env.api("admin1"), env.api("admin2")
    csv_data = "sku,warehouse_code,location_code,quantity,unit_cost\nART-A,BOD01,A-01,25,12.5\nART-X,BOD01,A-01,1,1\n"
    r = _upload(a, "opening_stock", csv_data)
    assert r.status_code == 422  # sin lote de saldos iniciales
    a.post("/imports/opening-batches", {"kind": "stock", "cutoff_date": str(TODAY), "source_description": "corto"},
           expect=422)
    batch = a.post("/imports/opening-batches", {"kind": "stock", "cutoff_date": str(TODAY),
                                                "source_description": "Inventario físico certificado al corte, "
                                                                      "exportado del sistema anterior"})
    job = _upload(a, "opening_stock", csv_data, opening_batch_id=batch["id"]).json()
    assert job["valid_rows"] == 1 and job["invalid_rows"] == 1
    job = a.post(f"/imports/jobs/{job['id']}/commit")
    assert scalar("select count(*) from stock_moves") == 0  # aún en staging
    a.post(f"/imports/opening-batches/{batch['id']}/approve", expect=403)  # quien crea no aprueba
    b.post(f"/imports/opening-batches/{batch['id']}/post", expect=422)  # sin aprobar
    b.post(f"/imports/opening-batches/{batch['id']}/approve")
    posted = b.post(f"/imports/opening-batches/{batch['id']}/post")
    assert posted["status"] == "posted"
    assert D(a.get(f"/inventory/availability/{env['items']['A']}")["available"]) == 25
    je = entries_for("opening_balance", batch["id"])
    assert len(je) == 1 and je[0].total_debit == D("312.50")
    assert scalar("select move_type from stock_moves") == "opening"


def test_opening_receivables(env):
    a, b = env.api("admin1"), env.api("admin2")
    batch = a.post("/imports/opening-batches", {"kind": "receivables", "cutoff_date": str(TODAY),
                                                "source_description": "Cartera por edades del sistema anterior"})
    csv_data = ("party_code,document_number,issue_date,due_date,currency,amount\n"
                f"CLI1,FV-OLD-1,{TODAY},{TODAY},COP,1000\nCLI1,FV-OLD-1,{TODAY},{TODAY},COP,5\n")
    job = _upload(a, "opening_receivables", csv_data, opening_batch_id=batch["id"]).json()
    assert job["valid_rows"] == 1  # duplicado dentro del archivo
    a.post(f"/imports/jobs/{job['id']}/commit")
    b.post(f"/imports/opening-batches/{batch['id']}/approve")
    b.post(f"/imports/opening-batches/{batch['id']}/post")
    ar = a.get("/accounting/open-items?kind=receivable")
    assert ar["total"] == 1 and D(ar["items"][0]["open_amount"]) == 1000
