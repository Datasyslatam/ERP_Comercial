"""Libros y reportes fiscales: Venezuela (libros de IVA, comprobante y TXT de retenciones) y Colombia
(auxiliar de IVA y certificados de retención)."""
import io
from datetime import timedelta
from decimal import Decimal

from conftest import PASSWORD, Api, build_company
from helpers import TODAY, D, confirmed_so, run_outbox, scalar, shipped_dispatch, stock_in
from sqlalchemy import select

from app.core.db import SessionLocal
from app.modules.accounting.fiscal import rif
from app.modules.accounting.models import Account, BankAccount
from app.modules.einvoicing.models import NumberingRange
from app.modules.iam.models import Role, UserCompanyRole
from app.modules.iam.service import create_user
from app.modules.inventory.models import Warehouse
from app.modules.masters.models import Item, Party, Tax, UnitOfMeasure
from app.modules.masters.validation import rif_check_digit


def _ve_company(client):
    with SessionLocal() as db:
        c = build_company(db, "VE", code="FISVE", special=True)
        role = db.scalar(select(Role).where(Role.code == "admin"))
        for n in ("fv1", "fv2"):
            u = create_user(db, email=f"{n}@pruebas.datasys.co", full_name=n, password=PASSWORD)
            db.add(UserCompanyRole(user_id=u.id, company_id=c.id, role_id=role.id))
        tax = {t.code: t for t in db.scalars(select(Tax).where(Tax.company_id == c.id))}
        und = db.scalar(select(UnitOfMeasure).where(UnitOfMeasure.code == "UND"))
        item = Item(company_id=c.id, sku="SRV-1", name="Servicio gravado", item_type="service", is_stocked=False,
                    base_uom_id=und.id, sales_tax_id=tax["IVA16"].id, list_price=Decimal("500"))
        cust = Party(company_id=c.id, code="CLI", legal_name="Cliente Especial C.A.", id_type="RIF", tax_id="J00124134",
                     check_digit=rif_check_digit("J00124134"), is_customer=True, is_withholding_agent=True,
                     address="Av. Bolívar, Valencia", email="cliente@pruebas.datasys.co", country_code="VE")
        supp = Party(company_id=c.id, code="PRV", legal_name="Proveedor de Servicios C.A.", id_type="RIF",
                     tax_id="J30000001", check_digit=rif_check_digit("J30000001"), is_supplier=True, country_code="VE",
                     email="proveedor@pruebas.datasys.co")
        supp.withholdings = [tax["RIVA75"]]
        db.add_all([item, cust, supp])
        gl = db.scalar(select(Account).where(Account.company_id == c.id, Account.code == "1.1.01.02"))
        bank = BankAccount(company_id=c.id, name="Banco Bs.", bank_name="Banco", account_number_masked="****1",
                           currency="VES", gl_account_id=gl.id)
        db.add(bank)
        for doc, prefix in (("invoice", "A"), ("credit_note", "NC")):
            db.add(NumberingRange(company_id=c.id, doc_type=doc, prefix=prefix, range_from=1, range_to=999,
                                  next_number=1, valid_from=TODAY - timedelta(days=1),
                                  valid_to=TODAY + timedelta(days=30), authorization_number="IMPRENTA-SBX",
                                  environment="sandbox"))
        expense = db.scalar(select(Account).where(Account.company_id == c.id, Account.code == "6.1.01"))
        wh = db.scalar(select(Warehouse).where(Warehouse.company_id == c.id))
        db.commit()
        ids = dict(c=c.id, item=item.id, cust=cust.id, supp=supp.id, bank=bank.id, wh=wh.id, site=wh.site_id,
                   expense=expense.id, **{k: v.id for k, v in tax.items()},
                   rif=rif("RIF", c.tax_id, c.tax_id_check_digit), supp_rif=rif("RIF", supp.tax_id, supp.check_digit))
    return ids, Api(client, "fv1@pruebas.datasys.co", ids["c"]), Api(client, "fv2@pruebas.datasys.co", ids["c"])


def _ve_sale(ids, a, b, qty=2):
    so = a.post("/sales/orders", {"site_id": ids["site"], "customer_id": ids["cust"], "warehouse_id": ids["wh"],
                                  "order_date": str(TODAY), "currency": "VES",
                                  "lines": [{"item_id": ids["item"], "quantity": str(qty)}]})
    b.post(f"/sales/orders/{so['id']}/confirm")
    d = a.post("/logistics/dispatches", {"sales_order_id": so["id"]})
    a.post(f"/logistics/dispatches/{d['id']}/start-picking")
    a.post(f"/logistics/dispatches/{d['id']}/picking", {"picks": [{"dispatch_line_id": d["lines"][0]["id"],
                                                                  "qty_picked": str(qty)}]})
    a.post(f"/logistics/dispatches/{d['id']}/pack", {"packages_count": 1})
    a.post(f"/logistics/dispatches/{d['id']}/ship", {})
    inv = a.post("/einvoicing/invoices/from-dispatch", {"dispatch_id": d["id"]})
    return a.post(f"/einvoicing/invoices/{inv['id']}/issue")


def _ve_purchase(ids, a, number, control=None, exempt=None):
    lines = [{"description": "Mantenimiento de equipos", "quantity": "1", "unit_price": "1000",
              "tax_id": ids["IVA16"], "expense_account_id": ids["expense"]}]
    if exempt:
        lines.append({"description": "Servicio exento", "quantity": "1", "unit_price": str(exempt),
                      "tax_id": ids["EXENTO"], "expense_account_id": ids["expense"]})
    body = {"supplier_id": ids["supp"], "supplier_invoice_number": number, "invoice_date": str(TODAY),
            "currency": "VES", "lines": lines}
    if control:
        body["supplier_control_number"] = control
    return a.post("/purchasing/invoices", body)


def test_venezuela_books_certificate_and_txt(client):
    ids, a, b = _ve_company(client)
    period = f"start={TODAY.replace(day=1)}&end={TODAY}"

    # Compras: sin número de control no se contabiliza (es obligatorio en el libro de compras)
    draft = _ve_purchase(ids, a, "F-100")
    r = a.post(f"/purchasing/invoices/{draft['id']}/post", expect=422)
    assert "número de control" in r.json()["error"]["message"]
    inv = _ve_purchase(ids, a, "F-101", control="00-000101", exempt=200)
    assert D(inv["withholding_total"]) == D("120.00")  # 75 % de un IVA de 160
    a.post(f"/purchasing/invoices/{inv['id']}/post")
    cert = scalar("select certificate_number from supplier_invoice_taxes where invoice_id = :i and is_withholding",
                  i=inv["id"])
    assert cert == f"{TODAY:%Y%m}00000001"

    # Ventas: factura, nota crédito y recaudo con el comprobante de retención que entrega el cliente
    sale = _ve_sale(ids, a, b)
    run_outbox()
    sale = a.get(f"/einvoicing/invoices/{sale['id']}")
    note = a.post(f"/einvoicing/invoices/{sale['id']}/notes", {
        "doc_type": "credit_note", "correction_code": "1", "reason": "Descuento por servicio incompleto",
        "lines": [{"invoice_line_id": sale["lines"][0]["id"], "quantity": "1"}]})
    note = a.post(f"/einvoicing/invoices/{note['id']}/issue")
    run_outbox()
    ar = next(i for i in a.get(f"/accounting/open-items?kind=receivable&party_id={ids['cust']}")["items"]
              if i["document_id"] == sale["id"])
    assert D(ar["open_amount"]) == D("580.00")
    a.post("/accounting/payments", {
        "kind": "receipt", "party_id": ids["cust"], "payment_date": str(TODAY), "bank_account_id": ids["bank"],
        "method": "transfer", "currency": "VES", "amount": "520", "idempotency_key": "cobro-fiscal-1",
        "withholdings": [{"tax_id": ids["RIVA75"], "base": "80", "amount": "60",
                          "certificate_number": "20261000000077"}],
        "allocations": [{"open_item_id": ar["id"], "amount": "580"}]})

    book = a.get(f"/fiscal/ve/sales-book?{period}")
    rows = {r["kind"]: r for r in book["rows"]}
    assert D(rows["invoice"]["total"]) == D("1160.00") and D(rows["invoice"]["general_tax"]) == D("160.00")
    assert rows["invoice"]["control_number"] and rows["invoice"]["taxpayer"] == "Contribuyente"
    assert D(rows["credit_note"]["general_base"]) == D("-500.00") and rows["credit_note"]["affected"] == sale["full_number"]
    assert rows["credit_note"]["credit_note"] == note["full_number"]
    assert rows["withholding"]["certificate"] == "20261000000077" and D(rows["withholding"]["withheld"]) == D("60.00")
    assert D(book["totals"]["general_tax"]) == D("80.00") and D(book["totals"]["withheld"]) == D("60.00")

    pbook = a.get(f"/fiscal/ve/purchase-book?{period}")
    assert len(pbook["rows"]) == 1  # el borrador F-100 no entra al libro
    row = pbook["rows"][0]
    assert (D(row["total"]), D(row["exempt"]), D(row["general_base"]), D(row["general_tax"]), D(row["withheld"])) == \
        (D("1360.00"), D("200.00"), D("1000.00"), D("160.00"), D("120.00"))
    assert row["certificate"] == cert and row["control_number"] == "00-000101" and not pbook["warnings"]

    # TXT de la quincena en curso para el portal del SENIAT
    half = 1 if TODAY.day <= 15 else 2
    r = a._call("GET", f"/fiscal/ve/vat-withholdings.txt?year={TODAY.year}&month={TODAY.month}&half={half}", 200)
    lines = r.text.strip().split("\r\n")
    assert len(lines) == 1
    f = lines[0].split("\t")
    assert len(f) == 16
    assert f[0] == ids["rif"] and f[1] == f"{TODAY:%Y%m}" and f[3:6] == ["C", "01", ids["supp_rif"]]
    assert f[6:11] == ["F-101", "00-000101", "1360.00", "1000.00", "120.00"]
    assert f[12:15] == [cert, "200.00", "16.00"]
    other = 2 if half == 1 else 1
    r = a._call("GET", f"/fiscal/ve/vat-withholdings.txt?year={TODAY.year}&month={TODAY.month}&half={other}", 200)
    assert r.text == ""

    # Comprobante en PDF y listado del periodo
    listed = a.get(f"/fiscal/ve/vat-withholdings?{period}")
    assert [x["certificate"] for x in listed] == [cert] and D(listed[0]["withheld"]) == D("120.00")
    pdf = a._call("GET", f"/fiscal/ve/vat-withholdings/{cert}/pdf", 200)
    assert pdf.content.startswith(b"%PDF") and cert in pdf.headers["content-disposition"]

    # Libro de ventas en Excel
    from openpyxl import load_workbook

    x = a._call("GET", f"/fiscal/ve/sales-book.xlsx?{period}", 200)
    ws = load_workbook(io.BytesIO(x.content)).active
    assert ws["A3"].value == "Libro de ventas" and ws["A1"].value.startswith("Empresa de Pruebas FISVE")

    # Retención al pagar: comparte la secuencia de comprobantes de IVA (consecutiva para el TXT)
    a.patch("/company/settings", {"settings": {"withholding_on_payment": True}})
    inv2 = _ve_purchase(ids, a, "F-102", control="00-000102")
    assert D(inv2["withholding_total"]) == 0
    a.post(f"/purchasing/invoices/{inv2['id']}/post")
    ap = next(i for i in a.get(f"/accounting/open-items?kind=payable&party_id={ids['supp']}")["items"]
              if i["document_id"] == inv2["id"])
    pay = a.post("/accounting/payments", {
        "kind": "disbursement", "party_id": ids["supp"], "payment_date": str(TODAY), "bank_account_id": ids["bank"],
        "method": "transfer", "currency": "VES", "amount": "1040", "idempotency_key": "pago-fiscal-1",
        "withholdings": [{"tax_id": ids["RIVA75"], "base": "160", "amount": "120"}],
        "allocations": [{"open_item_id": ap["id"], "amount": "1160"}]})
    assert pay["withholdings"][0]["certificate_number"] == f"{TODAY:%Y%m}00000002"
    assert len(a.get(f"/fiscal/ve/vat-withholdings?{period}")) == 2

    # Los reportes de Colombia no aplican a una empresa venezolana
    a.get(f"/fiscal/co/vat-auxiliary?{period}", expect=422)


def test_colombia_vat_auxiliary_and_certificates(env):
    from test_treasury import _supplier_invoice

    a = env.api("contador")
    _, item = _supplier_invoice(env, number="FE-3001")  # base 1.000, IVA 190, retención 2,5 % = 25
    stock_in(env, "A", qty=10, price=100)
    so = confirmed_so(env, "A", qty=2, price=200)
    d = shipped_dispatch(env, so)
    adm = env.api("admin1")
    inv = adm.post("/einvoicing/invoices/from-dispatch", {"dispatch_id": d["id"]})
    adm.post(f"/einvoicing/invoices/{inv['id']}/issue")
    period = f"start={TODAY.replace(day=1)}&end={TODAY}"

    aux = a.get(f"/fiscal/co/vat-auxiliary?{period}")
    assert D(aux["generated"]["total"]) == D("76.00")  # 19 % de 400
    assert D(aux["deductible"]["total"]) == D("190.00")  # factura de proveedor de 1.000 (la recepción sin factura no cuenta)
    assert D(aux["summary"]["balance"]) == D("-114.00")
    assert aux["generated"]["by_rate"][0]["rate"].startswith("19")

    certs = a.get(f"/fiscal/co/withholding-certificates?{period}&kind=income")
    assert len(certs) == 1 and certs[0]["party_id"] == env["supplier_id"]
    assert D(certs[0]["amount"]) == D("25.00") and D(certs[0]["base"]) == D("1000.00")
    assert certs[0]["lines"][0]["code"] == "RTEC25"
    assert a.get(f"/fiscal/co/withholding-certificates?{period}&kind=vat") == []
    pdf = a._call("GET", f"/fiscal/co/withholding-certificates/{env['supplier_id']}/pdf?{period}&kind=income", 200)
    assert pdf.content.startswith(b"%PDF")
    a._call("GET", f"/fiscal/co/withholding-certificates/{env['supplier_id']}/pdf?{period}&kind=ica", 404)
    x = a._call("GET", f"/fiscal/co/vat-auxiliary.xlsx?{period}", 200)
    assert x.content[:2] == b"PK"
    # Permisos: el vendedor no consulta reportes fiscales; los libros de Venezuela no aplican en Colombia
    env.api("vendedor").get(f"/fiscal/co/vat-auxiliary?{period}", expect=403)
    a.get(f"/fiscal/ve/sales-book?{period}", expect=422)
    assert scalar("select count(*) from audit_log where action = 'export' and entity_type = 'withholding_certificate'") \
        == 1
