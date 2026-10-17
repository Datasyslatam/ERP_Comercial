"""Importación de datos: plantilla → carga → mapeo → validación → vista previa → corrección → importación por filas
(con savepoint) → conciliación. Los saldos iniciales se cargan en un lote con fecha de corte, origen documentado
y aprobación de un usuario distinto a quien lo creó; solo entonces se contabilizan."""
import csv
import hashlib
import io
from collections import defaultdict
from datetime import UTC, date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.audit import log_action
from app.core.crud import get_owned
from app.core.deps import Principal
from app.core.errors import BusinessRuleViolation, Conflict
from app.core.money import ZERO, money
from app.core.policies import enforce_sod
from app.modules.accounting.ledger import LineSpec, post_entry, posting_account
from app.modules.accounting.models import Account, OpenItem
from app.modules.imports.models import ImportJob, ImportRow, OpeningBalanceBatch
from app.modules.inventory import stock
from app.modules.inventory.models import Location, Warehouse
from app.modules.masters.models import Item, ItemCategory, Party, PaymentTerm, Tax, UnitOfMeasure
from app.modules.masters.validation import validate_tax_id

# campo: (requerido, tipo, descripción)
SPECS: dict[str, dict[str, tuple[bool, str, str]]] = {
    "parties": {
        "code": (True, "str", "Código interno único"), "legal_name": (True, "str", "Razón social o nombre"),
        "id_type": (True, "str", "NIT, CC, CE, PAS, RIF, CI, EXT"), "tax_id": (True, "str", "Número sin DV"),
        "check_digit": (False, "str", "Dígito de verificación (NIT/RIF)"),
        "is_customer": (False, "bool", "SI/NO"), "is_supplier": (False, "bool", "SI/NO"),
        "is_carrier": (False, "bool", "SI/NO"), "email": (False, "str", "Correo de facturación"),
        "phone": (False, "str", ""), "address": (False, "str", ""), "city": (False, "str", ""),
        "country_code": (False, "str", "CO/VE"), "payment_term_code": (False, "str", "Código condición de pago"),
        "credit_limit": (False, "decimal", "Cupo en moneda funcional"),
    },
    "items": {
        "sku": (True, "str", "Código único"), "name": (True, "str", "Descripción"),
        "item_type": (True, "str", "raw_material, supply, finished_good, merchandise, service"),
        "base_uom_code": (True, "str", "UND, KG, L..."), "category_code": (False, "str", ""),
        "is_stocked": (False, "bool", "SI/NO (por defecto SI salvo servicios)"),
        "track_lots": (False, "bool", ""), "track_expiry": (False, "bool", ""),
        "requires_inspection": (False, "bool", ""), "sales_tax_code": (False, "str", "Ej. IVA19 / IVA16"),
        "purchase_tax_code": (False, "str", ""), "min_stock": (False, "decimal", ""),
        "max_stock": (False, "decimal", ""), "list_price": (False, "decimal", ""), "barcode": (False, "str", ""),
    },
    "accounts": {
        "code": (True, "str", ""), "name": (True, "str", ""),
        "account_type": (True, "str", "asset, liability, equity, revenue, expense, cost"),
        "parent_code": (False, "str", ""), "is_postable": (False, "bool", ""),
        "requires_party": (False, "bool", ""), "is_cash": (False, "bool", ""),
    },
    "opening_stock": {
        "sku": (True, "str", ""), "warehouse_code": (True, "str", ""), "location_code": (True, "str", ""),
        "quantity": (True, "decimal", "> 0"), "unit_cost": (True, "decimal", "Costo unitario en moneda funcional"),
        "lot_number": (False, "str", ""), "expiry_date": (False, "date", "AAAA-MM-DD"),
    },
    "opening_receivables": {
        "party_code": (True, "str", "Código del cliente"), "document_number": (True, "str", ""),
        "issue_date": (True, "date", ""), "due_date": (True, "date", ""), "currency": (True, "str", ""),
        "exchange_rate": (False, "decimal", "Obligatoria si la moneda no es la funcional"),
        "amount": (True, "decimal", "Saldo pendiente a la fecha de corte"),
    },
}
SPECS["opening_payables"] = {**SPECS["opening_receivables"], "party_code": (True, "str", "Código del proveedor")}
OPENING_KIND = {"opening_stock": "stock", "opening_receivables": "receivables", "opening_payables": "payables"}
_TRUE = {"si", "sí", "s", "true", "1", "x", "yes", "y"}


def template(entity: str, fmt: str = "xlsx") -> tuple[bytes, str]:
    spec = SPECS[entity]
    if fmt == "csv":
        buf = io.StringIO()
        csv.writer(buf).writerow(spec.keys())
        return buf.getvalue().encode("utf-8-sig"), "text/csv"
    wb = Workbook()
    ws = wb.active
    ws.title = "datos"
    ws.append(list(spec.keys()))
    info = wb.create_sheet("instrucciones")
    info.append(["campo", "obligatorio", "tipo", "descripción"])
    for k, (req, typ, desc) in spec.items():
        info.append([k, "SI" if req else "NO", typ, desc])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _parse(filename: str, data: bytes) -> list[dict]:
    if filename.lower().endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(h).strip() if h is not None else "" for h in rows[0]]
        return [{headers[i]: v for i, v in enumerate(r) if i < len(headers) and headers[i]}
                for r in rows[1:] if any(v not in (None, "") for v in r)]
    text_data = data.decode("utf-8-sig")
    dialect = csv.Sniffer().sniff(text_data.splitlines()[0], delimiters=",;\t") if text_data else csv.excel
    return [r for r in csv.DictReader(io.StringIO(text_data), dialect=dialect) if any((v or "").strip() for v in r.values())]


def upload(db: Session, p: Principal, entity: str, filename: str, data: bytes, mapping: dict | None,
           opening_batch_id: int | None) -> ImportJob:
    if entity not in SPECS:
        raise BusinessRuleViolation(f"Entidad no importable: {entity}")
    if entity in OPENING_KIND:
        if not opening_batch_id:
            raise BusinessRuleViolation("Los saldos iniciales requieren un lote de saldos iniciales (fecha de corte)")
        batch = get_owned(db, OpeningBalanceBatch, opening_batch_id, p.company_id)
        if batch.kind != OPENING_KIND[entity] or batch.status != "draft":
            raise BusinessRuleViolation("El lote no corresponde al tipo de saldo o ya no está en borrador")
    rows = _parse(filename, data)
    if not rows:
        raise BusinessRuleViolation("El archivo no contiene filas de datos")
    if len(rows) > 50000:
        raise BusinessRuleViolation("Máximo 50.000 filas por archivo; divida la carga en lotes")
    sha = hashlib.sha256(data).hexdigest()
    mapping = mapping or {}
    unknown = set(mapping.values()) - set(SPECS[entity])
    if unknown:
        raise BusinessRuleViolation(f"El mapeo apunta a campos inexistentes: {sorted(unknown)}")
    job = ImportJob(company_id=p.company_id, entity=entity, filename=filename[:255], file_sha256=sha,
                    column_mapping=mapping, uploaded_by_id=p.user_id, total_rows=len(rows),
                    opening_batch_id=opening_batch_id)
    for i, raw in enumerate(rows, start=2):
        mapped = {}
        for col, val in raw.items():
            field = mapping.get(col, col)
            if field in SPECS[entity]:
                mapped[field] = val.isoformat() if isinstance(val, date | datetime) else (
                    str(val).strip() if val is not None else None)
        job.rows.append(ImportRow(row_number=i, data=mapped))
    db.add(job)
    db.flush()
    previous = db.scalar(select(ImportJob.id).where(ImportJob.company_id == p.company_id, ImportJob.file_sha256 == sha,
                                                    ImportJob.id != job.id, ImportJob.status.in_(
                                                        ("imported", "partially_imported"))))
    log_action(db, "import", "import_jobs", job.id, f"{entity}: {filename} ({len(rows)} filas)"
               + (f"; archivo idéntico ya importado en lote {previous}" if previous else ""))
    return job


def _coerce(value, typ: str):
    if value in (None, ""):
        return None
    if typ == "decimal":
        try:
            return Decimal(str(value).replace(",", ".")) if str(value).count(",") == 1 and "." not in str(value) \
                else Decimal(str(value).replace(",", ""))
        except InvalidOperation as exc:
            raise ValueError("número inválido") from exc
    if typ == "bool":
        return str(value).strip().lower() in _TRUE
    if typ == "date":
        try:
            return date.fromisoformat(str(value)[:10])
        except ValueError as exc:
            raise ValueError("fecha inválida (AAAA-MM-DD)") from exc
    return str(value).strip()


class _Ctx:
    """Caché de referencias para validar en bloque sin una consulta por fila."""

    def __init__(self, db: Session, company_id: int):
        c = company_id
        self.uoms = {u.code: u.id for u in db.scalars(select(UnitOfMeasure))}
        self.cats = {x.code: x.id for x in db.scalars(select(ItemCategory).where(ItemCategory.company_id == c))}
        self.taxes = {x.code: x for x in db.scalars(select(Tax).where(Tax.company_id == c))}
        self.terms = {x.code: x.id for x in db.scalars(select(PaymentTerm).where(PaymentTerm.company_id == c))}
        self.items = {x.sku: x for x in db.scalars(select(Item).where(Item.company_id == c))}
        self.parties = {x.code: x for x in db.scalars(select(Party).where(Party.company_id == c))}
        self.party_ids = {(x.id_type, x.tax_id) for x in self.parties.values()}
        self.accounts = {x.code: x for x in db.scalars(select(Account).where(Account.company_id == c))}
        self.whs = {x.code: x.id for x in db.scalars(select(Warehouse).where(Warehouse.company_id == c))}
        self.locs = {(x.warehouse_id, x.code): x.id for x in db.scalars(select(Location).where(Location.company_id == c))}
        self.docs = {(x.kind, x.party_id, x.document_number) for x in db.scalars(
            select(OpenItem).where(OpenItem.company_id == c))}


def _validate_row(entity: str, data: dict, ctx: _Ctx, seen: set, company) -> tuple[dict, list[str]]:
    errors, v = [], {}
    for field, (req, typ, _) in SPECS[entity].items():
        try:
            v[field] = _coerce(data.get(field), typ)
        except ValueError as exc:
            errors.append(f"{field}: {exc}")
            continue
        if req and v[field] in (None, ""):
            errors.append(f"{field}: obligatorio")
    if errors:
        return v, errors
    if entity == "parties":
        key = (v["id_type"].upper(), v["tax_id"])
        if v["code"] in ctx.parties or key in ctx.party_ids:
            errors.append("Duplicado: el tercero ya existe (código o identificación)")
        if ("p", v["code"]) in seen or ("pid",) + key in seen:
            errors.append("Duplicado dentro del archivo")
        seen.update({("p", v["code"]), ("pid",) + key})
        try:
            validate_tax_id(v["id_type"], v["tax_id"], v.get("check_digit"), v.get("country_code") or "CO")
        except BusinessRuleViolation as exc:
            errors.append(exc.message)
        if not (v.get("is_customer") or v.get("is_supplier") or v.get("is_carrier")):
            errors.append("Debe marcarse como cliente, proveedor o transportista")
        if v.get("payment_term_code") and v["payment_term_code"] not in ctx.terms:
            errors.append("Condición de pago inexistente")
    elif entity == "items":
        if v["sku"] in ctx.items or ("i", v["sku"]) in seen:
            errors.append("Duplicado: el SKU ya existe")
        seen.add(("i", v["sku"]))
        if v["base_uom_code"] not in ctx.uoms:
            errors.append("Unidad de medida inexistente")
        if v.get("category_code") and v["category_code"] not in ctx.cats:
            errors.append("Categoría inexistente")
        for f in ("sales_tax_code", "purchase_tax_code"):
            if v.get(f) and v[f] not in ctx.taxes:
                errors.append(f"{f}: impuesto inexistente")
        if v["item_type"] not in ("raw_material", "supply", "wip", "finished_good", "merchandise", "service"):
            errors.append("Tipo de artículo inválido")
    elif entity == "accounts":
        if v["code"] in ctx.accounts or ("a", v["code"]) in seen:
            errors.append("Duplicado: la cuenta ya existe")
        seen.add(("a", v["code"]))
        if v["account_type"] not in ("asset", "liability", "equity", "revenue", "expense", "cost"):
            errors.append("Tipo de cuenta inválido")
        if v.get("parent_code") and v["parent_code"] not in ctx.accounts and ("a", v["parent_code"]) not in seen:
            errors.append("Cuenta padre inexistente (debe existir o venir antes en el archivo)")
    elif entity == "opening_stock":
        item = ctx.items.get(v["sku"])
        wh = ctx.whs.get(v["warehouse_code"])
        if item is None:
            errors.append("Artículo inexistente")
        elif not item.is_stocked:
            errors.append("El artículo no es inventariable")
        elif item.track_lots and not v.get("lot_number"):
            errors.append("El artículo exige lote")
        elif item.track_expiry and not v.get("expiry_date"):
            errors.append("El artículo exige fecha de vencimiento")
        if wh is None or (wh, v["location_code"]) not in ctx.locs:
            errors.append("Bodega/ubicación inexistente")
        if v["quantity"] is not None and v["quantity"] <= 0:
            errors.append("La cantidad debe ser positiva")
        if v["unit_cost"] is not None and v["unit_cost"] < 0:
            errors.append("El costo no puede ser negativo")
        key = ("s", v["sku"], v["warehouse_code"], v["location_code"], v.get("lot_number"))
        if key in seen:
            errors.append("Duplicado dentro del archivo (misma ubicación y lote)")
        seen.add(key)
    else:  # cartera inicial
        kind = "receivable" if entity == "opening_receivables" else "payable"
        party = ctx.parties.get(v["party_code"])
        if party is None:
            errors.append("Tercero inexistente")
        elif (kind == "receivable" and not party.is_customer) or (kind == "payable" and not party.is_supplier):
            errors.append("El tercero no tiene el rol requerido")
        elif (kind, party.id, v["document_number"]) in ctx.docs or (kind, party.id, v["document_number"]) in seen:
            errors.append("Duplicado: documento ya cargado para el tercero")
        else:
            seen.add((kind, party.id, v["document_number"]))
        if v["amount"] is not None and v["amount"] <= 0:
            errors.append("El saldo debe ser positivo")
        if v["currency"] != company.functional_currency and not v.get("exchange_rate"):
            errors.append("Falta la tasa de cambio para moneda extranjera")
        if v["due_date"] and v["issue_date"] and v["due_date"] < v["issue_date"]:
            errors.append("El vencimiento es anterior a la emisión")
    return v, errors


def validate(db: Session, p: Principal, job_id: int) -> ImportJob:
    job = get_owned(db, ImportJob, job_id, p.company_id, lock=True)
    if job.status in ("imported", "cancelled"):
        raise Conflict("El lote ya fue importado o cancelado")
    ctx, seen = _Ctx(db, p.company_id), set()
    valid = invalid = 0
    for row in job.rows:
        if row.status == "imported":
            continue
        _, errors = _validate_row(job.entity, row.data, ctx, seen, p.company)
        row.status, row.errors = ("invalid", errors) if errors else ("valid", [])
        valid, invalid = valid + (not errors), invalid + bool(errors)
    job.valid_rows, job.invalid_rows = valid, invalid
    if job.status == "uploaded":
        job.status = "validated"
    return job


def update_row(db: Session, p: Principal, row_id: int, data: dict) -> ImportRow:
    row = db.get(ImportRow, row_id)
    job = get_owned(db, ImportJob, row.job_id, p.company_id) if row else None
    if job is None:
        raise BusinessRuleViolation("Fila inexistente")
    if row.status == "imported":
        raise Conflict("La fila ya fue importada")
    unknown = set(data) - set(SPECS[job.entity])
    if unknown:
        raise BusinessRuleViolation(f"Campos desconocidos: {sorted(unknown)}")
    row.data = {**row.data, **{k: (str(v) if v is not None else None) for k, v in data.items()}}
    row.status, row.errors = "pending", []
    return row


def _create(db: Session, p: Principal, job: ImportJob, v: dict, ctx: _Ctx):
    c = p.company_id
    if job.entity == "parties":
        obj = Party(company_id=c, code=v["code"], legal_name=v["legal_name"], id_type=v["id_type"].upper(),
                    tax_id=v["tax_id"], check_digit=v.get("check_digit"), is_customer=bool(v.get("is_customer")),
                    is_supplier=bool(v.get("is_supplier")), is_carrier=bool(v.get("is_carrier")), email=v.get("email"),
                    phone=v.get("phone"), address=v.get("address"), city=v.get("city"),
                    country_code=v.get("country_code") or p.company.country_code,
                    payment_term_id=ctx.terms.get(v.get("payment_term_code")),
                    credit_limit=v.get("credit_limit") or ZERO)
        ctx.parties[obj.code] = obj
    elif job.entity == "items":
        is_service = v["item_type"] == "service"
        obj = Item(company_id=c, sku=v["sku"], name=v["name"], item_type=v["item_type"],
                   base_uom_id=ctx.uoms[v["base_uom_code"]], category_id=ctx.cats.get(v.get("category_code")),
                   is_stocked=False if is_service else (v.get("is_stocked") if v.get("is_stocked") is not None else True),
                   track_lots=bool(v.get("track_lots")), track_expiry=bool(v.get("track_expiry")),
                   requires_inspection=bool(v.get("requires_inspection")),
                   sales_tax_id=ctx.taxes[v["sales_tax_code"]].id if v.get("sales_tax_code") else None,
                   purchase_tax_id=ctx.taxes[v["purchase_tax_code"]].id if v.get("purchase_tax_code") else None,
                   min_stock=v.get("min_stock") or ZERO, max_stock=v.get("max_stock") or ZERO,
                   list_price=v.get("list_price") or ZERO, barcode=v.get("barcode"))
        ctx.items[obj.sku] = obj
    elif job.entity == "accounts":
        parent = ctx.accounts.get(v.get("parent_code")) if v.get("parent_code") else None
        obj = Account(company_id=c, code=v["code"], name=v["name"], account_type=v["account_type"],
                      normal_balance="debit" if v["account_type"] in ("asset", "expense", "cost") else "credit",
                      parent_id=parent.id if parent else None,
                      is_postable=v.get("is_postable") if v.get("is_postable") is not None else True,
                      requires_party=bool(v.get("requires_party")), is_cash=bool(v.get("is_cash")))
        db.add(obj)
        db.flush()
        ctx.accounts[obj.code] = obj
        return obj
    else:
        return None  # los saldos iniciales quedan en staging hasta aprobar el lote
    db.add(obj)
    db.flush()
    return obj


def commit(db: Session, p: Principal, job_id: int) -> ImportJob:
    job = validate(db, p, job_id)
    ctx = _Ctx(db, p.company_id)
    imported = 0
    for row in job.rows:
        if row.status != "valid":
            continue
        v, errors = _validate_row(job.entity, row.data, ctx, set(), p.company)
        if errors:
            row.status, row.errors = "invalid", errors
            continue
        try:
            with db.begin_nested():
                obj = _create(db, p, job, v, ctx)
            row.status, row.created_entity_id = "imported", getattr(obj, "id", None)
            imported += 1
        except Exception as exc:  # noqa: BLE001 - la fila se rechaza y el lote continúa
            row.status, row.errors = "rejected", [f"{type(exc).__name__}: {str(exc)[:300]}"]
    db.flush()
    job.imported_rows = sum(1 for r in job.rows if r.status == "imported")
    job.invalid_rows = sum(1 for r in job.rows if r.status in ("invalid", "rejected"))
    job.status = "imported" if job.imported_rows == job.total_rows else (
        "partially_imported" if job.imported_rows else "failed")
    job.finished_at = datetime.now(UTC)
    log_action(db, "import", "import_jobs", job.id, f"Importadas {job.imported_rows}/{job.total_rows} filas")
    return job


# ---------------------------------------------------------------- saldos iniciales
def create_opening_batch(db: Session, p: Principal, data: dict) -> OpeningBalanceBatch:
    if len(data.get("source_description") or "") < 10:
        raise BusinessRuleViolation("Documente el origen de los saldos (sistema anterior, informe, fecha de extracción)")
    if data["cutoff_date"] > date.today():
        raise BusinessRuleViolation("La fecha de corte no puede ser futura")
    b = OpeningBalanceBatch(company_id=p.company_id, created_by_id=p.user_id, **data)
    db.add(b)
    db.flush()
    return b


def _batch_rows(db: Session, batch: OpeningBalanceBatch) -> list[tuple[ImportJob, ImportRow]]:
    jobs = db.scalars(select(ImportJob).where(ImportJob.opening_batch_id == batch.id)).all()
    return [(j, r) for j in jobs for r in j.rows if r.status == "imported"]


def batch_summary(db: Session, p: Principal, batch_id: int) -> dict:
    batch = get_owned(db, OpeningBalanceBatch, batch_id, p.company_id)
    total, n = ZERO, 0
    for job, row in _batch_rows(db, batch):
        d = row.data
        if job.entity == "opening_stock":
            total += money(Decimal(d["quantity"]) * Decimal(d["unit_cost"]))
        else:
            rate = Decimal(d.get("exchange_rate") or 1)
            total += money(Decimal(d["amount"]) * rate)
        n += 1
    return {"batch_id": batch.id, "kind": batch.kind, "status": batch.status, "cutoff_date": batch.cutoff_date,
            "rows": n, "total_functional": total, "journal_entry_id": batch.journal_entry_id}


def approve_batch(db: Session, p: Principal, batch_id: int) -> OpeningBalanceBatch:
    b = get_owned(db, OpeningBalanceBatch, batch_id, p.company_id, lock=True)
    if b.status != "draft":
        raise Conflict("El lote no está en borrador")
    enforce_sod(p, "sod_opening_balance_self_approval", b.created_by_id, "aprobar")
    if not _batch_rows(db, b):
        raise BusinessRuleViolation("El lote no tiene filas importadas")
    b.status, b.approved_by_id, b.approved_at = "approved", p.user_id, datetime.now(UTC)
    log_action(db, "approve", "opening_balance_batches", b.id, f"Corte {b.cutoff_date}")
    return b


def post_batch(db: Session, p: Principal, batch_id: int) -> OpeningBalanceBatch:
    b = get_owned(db, OpeningBalanceBatch, batch_id, p.company_id, lock=True)
    if b.status == "posted":
        return b
    if b.status != "approved":
        raise BusinessRuleViolation("El lote debe estar aprobado")
    ctx = _Ctx(db, p.company_id)
    opening = posting_account(db, p.company_id, "opening_balance")
    lines: list[LineSpec] = []
    by_item: dict[int, Decimal] = defaultdict(lambda: ZERO)
    when = datetime.combine(b.cutoff_date, datetime.min.time(), tzinfo=UTC)
    for job, row in _batch_rows(db, b):
        d = row.data
        if job.entity == "opening_stock":
            item = ctx.items[d["sku"]]
            loc_id = ctx.locs[(ctx.whs[d["warehouse_code"]], d["location_code"])]
            lot = stock.get_or_create_lot(db, p.company_id, item, d.get("lot_number"),
                                          expiry_date=date.fromisoformat(d["expiry_date"]) if d.get("expiry_date") else None)
            mv = stock.receive(db, p.company_id, item_id=item.id, location_id=loc_id, quantity=Decimal(d["quantity"]),
                               unit_cost=Decimal(d["unit_cost"]), move_type="opening", lot_id=lot.id if lot else None,
                               src=stock.MoveSource("opening_balance", b.id, row.id, f"SI-{b.id}"),
                               reason=f"Saldo inicial al {b.cutoff_date}", user_id=p.user_id, move_date=when)
            by_item[item.id] += mv.total_cost
        else:
            kind = "receivable" if job.entity == "opening_receivables" else "payable"
            party = ctx.parties[d["party_code"]]
            rate = Decimal(d.get("exchange_rate") or 1)
            amount = money(Decimal(d["amount"]))
            oi = OpenItem(company_id=p.company_id, kind=kind, party_id=party.id, document_type="opening",
                          document_id=row.id, document_number=d["document_number"],
                          issue_date=date.fromisoformat(d["issue_date"]), due_date=date.fromisoformat(d["due_date"]),
                          currency=d["currency"], exchange_rate=rate, original_amount=amount, open_amount=amount)
            db.add(oi)
            value = money(amount * rate)
            acc = posting_account(db, p.company_id, "accounts_receivable" if kind == "receivable" else "accounts_payable")
            lines.append(LineSpec(account_id=acc, debit=value if kind == "receivable" else ZERO,
                                  credit=value if kind == "payable" else ZERO, party_id=party.id,
                                  currency=d["currency"], amount_currency=amount, exchange_rate=rate,
                                  description=f"Saldo inicial {d['document_number']}"))
    if by_item:
        from app.modules.accounting.posting import _item_account

        inv: dict[int, Decimal] = defaultdict(lambda: ZERO)
        for item_id, value in by_item.items():
            inv[_item_account(db, p.company_id, item_id, "inventory_account_id", "inventory")] += value
        lines += [LineSpec(account_id=a, debit=money(v), description="Inventario inicial") for a, v in inv.items()]
    net = sum((ln.debit - ln.credit for ln in lines), ZERO)
    lines.append(LineSpec(account_id=opening, debit=-net if net < 0 else ZERO, credit=net if net > 0 else ZERO,
                          description="Contrapartida saldos iniciales (cuenta puente)"))
    entry = post_entry(db, company_id=p.company_id, entry_date=b.cutoff_date, journal_type="opening",
                       description=f"Saldos iniciales {b.kind} al {b.cutoff_date}: {b.source_description[:200]}",
                       lines=lines, idempotency_key=f"opening_batch:{b.id}", source_type="opening_balance",
                       source_id=b.id, user_id=p.user_id)
    for oi in db.scalars(select(OpenItem).where(OpenItem.document_type == "opening", OpenItem.journal_entry_id.is_(None),
                                                OpenItem.company_id == p.company_id)):
        oi.journal_entry_id = entry.id
    b.status, b.posted_at, b.journal_entry_id = "posted", datetime.now(UTC), entry.id
    log_action(db, "post", "opening_balance_batches", b.id, f"Asiento {entry.number}")
    return b


def reconcile_job(db: Session, p: Principal, job_id: int) -> dict:
    """Conciliación posterior: compara lo cargado contra lo persistido en el sistema."""
    job = get_owned(db, ImportJob, job_id, p.company_id)
    imported = [r for r in job.rows if r.status == "imported"]
    report = {"job_id": job.id, "entity": job.entity, "file_rows": job.total_rows, "imported_rows": len(imported),
              "rejected_rows": sum(1 for r in job.rows if r.status in ("invalid", "rejected"))}
    if job.entity in ("parties", "items", "accounts"):
        model = {"parties": Party, "items": Item, "accounts": Account}[job.entity]
        ids = [r.created_entity_id for r in imported if r.created_entity_id]
        found = db.scalar(select(func.count()).select_from(model).where(model.id.in_(ids))) if ids else 0
        report.update({"persisted": found, "ok": found == len(imported)})
    else:
        batch = db.get(OpeningBalanceBatch, job.opening_batch_id)
        report.update(batch_summary(db, p, batch.id))
        if batch.status == "posted" and batch.journal_entry_id:
            from app.modules.accounting.models import JournalEntry

            je = db.get(JournalEntry, batch.journal_entry_id)
            report["journal_total"] = je.total_debit
    return report
