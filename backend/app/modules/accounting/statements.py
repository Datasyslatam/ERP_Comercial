"""Lectura de extractos bancarios desde archivo: CSV, Excel (.xlsx) y OFX/QFX.

Los bancos de Colombia y Venezuela exportan formatos distintos; se reconocen los encabezados más comunes
(fecha, descripción/concepto, referencia/documento, valor/monto o débito y crédito por separado, saldo), fechas
día/mes/año o ISO y números en formato local (1.234,56) o internacional (1,234.56). El resultado es una vista previa:
el usuario revisa y confirma antes de crear el extracto.
"""
import csv
import io
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from app.core.errors import BusinessRuleViolation

MAX_BYTES = 5 * 1024 * 1024

# El orden importa: débito y crédito se reconocen antes que "valor" ("Valor débito" no es la columna de valor).
HEADERS = {
    "date": ("fecha", "date", "fecha transaccion", "fecha movimiento", "fecha operacion", "fecha valor"),
    "debit": ("debito", "debitos", "cargo", "cargos", "retiro", "retiros", "debit", "egreso", "salida"),
    "credit": ("credito", "creditos", "abono", "abonos", "deposito", "depositos", "credit", "ingreso", "entrada"),
    "balance": ("saldo", "balance"),
    "description": ("descripcion", "concepto", "detalle", "description", "movimiento", "transaccion", "observacion"),
    "reference": ("referencia", "documento", "ref", "reference", "numero", "comprobante", "serial"),
    "amount": ("valor", "monto", "importe", "amount"),
}


@dataclass
class ParsedStatement:
    format: str
    lines: list[dict] = field(default_factory=list)
    opening_balance: Decimal | None = None
    closing_balance: Decimal | None = None
    statement_date: date | None = None
    warnings: list[str] = field(default_factory=list)

    def as_dict(self) -> dict:
        total = sum((Decimal(ln["amount"]) for ln in self.lines), Decimal("0"))
        return {"format": self.format, "lines": self.lines, "opening_balance": self.opening_balance,
                "closing_balance": self.closing_balance, "statement_date": self.statement_date,
                "net_movement": total, "warnings": self.warnings}


def _norm(s: str) -> str:
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().lower().strip()
    return re.sub(r"\s+", " ", s)


def parse_number(raw) -> Decimal | None:
    """1.234,56 · 1,234.56 · -850000 · (1.200,00) · $ 1.200 → Decimal."""
    if raw is None:
        return None
    if isinstance(raw, int | float | Decimal):
        return Decimal(str(raw))
    s = str(raw).strip()
    if not s:
        return None
    negative = s.startswith("(") and s.endswith(")") or s.startswith("-") or s.endswith("-")
    s = re.sub(r"[^\d.,]", "", s)
    if not s:
        return None
    if "," in s and "." in s:
        decimal_sep = "," if s.rfind(",") > s.rfind(".") else "."
        thousands = "." if decimal_sep == "," else ","
        s = s.replace(thousands, "").replace(decimal_sep, ".")
    elif "," in s:
        parts = s.split(",")
        s = s.replace(",", ".") if len(parts) == 2 and len(parts[1]) in (1, 2) else s.replace(",", "")
    elif s.count(".") > 1 or (s.count(".") == 1 and len(s.split(".")[1]) == 3):
        s = s.replace(".", "")  # 1.234.567 o 1.234 (miles)
    try:
        value = Decimal(s)
    except InvalidOperation:
        return None
    return -value if negative else value


def parse_date(raw) -> date | None:
    if raw is None or raw == "":
        return None
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    s = str(raw).strip()[:19]
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y/%m/%d", "%d/%m/%y", "%Y%m%d", "%Y-%m-%d %H:%M:%S",
                "%d/%m/%Y %H:%M:%S"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _map_headers(row: list) -> dict[str, int] | None:
    names = [_norm(c) for c in row]
    found: dict[str, int] = {}
    for key, options in HEADERS.items():
        for i, n in enumerate(names):
            if i in found.values():
                continue
            if n in options or any(o in n for o in options if len(o) >= 5):
                found[key] = i
                break
    if "date" in found and ("amount" in found or "debit" in found or "credit" in found):
        return found
    return None


def _rows_to_statement(rows: list[list], fmt: str) -> ParsedStatement:
    st = ParsedStatement(format=fmt)
    cols = None
    start = 0
    for i, row in enumerate(rows[:30]):  # el encabezado puede estar después de los datos del banco
        cols = _map_headers(row)
        if cols:
            start = i + 1
            break
    if cols is None:
        raise BusinessRuleViolation("No se reconocieron los encabezados del extracto (se requieren fecha y valor, "
                                    "o débito y crédito)")
    first_balance = None

    def cell(row: list, key: str):
        return row[cols[key]] if key in cols and cols[key] < len(row) else None

    for n, row in enumerate(rows[start:], start=start + 1):
        d = parse_date(cell(row, "date"))
        if d is None:
            if any(str(c).strip() for c in row if c is not None):
                st.warnings.append(f"Fila {n} omitida: fecha no reconocida")
            continue
        if "amount" in cols:
            amount = parse_number(cell(row, "amount"))
        else:
            debit, credit = parse_number(cell(row, "debit")) or Decimal("0"), parse_number(cell(row, "credit")) or Decimal("0")
            amount = abs(credit) - abs(debit)
        if amount is None or amount == 0:
            st.warnings.append(f"Fila {n} omitida: sin valor")
            continue
        description = str(cell(row, "description") or "").strip()[:255] or "Movimiento bancario"
        reference = str(cell(row, "reference") or "").strip()[:80] or None
        st.lines.append({"txn_date": d, "description": description, "reference": reference, "amount": amount})
        bal = parse_number(cell(row, "balance"))
        if bal is not None:
            if first_balance is None:
                first_balance = bal - amount
            st.closing_balance = bal
    if not st.lines:
        raise BusinessRuleViolation("El archivo no contiene movimientos reconocibles")
    st.opening_balance = first_balance
    st.statement_date = max(ln["txn_date"] for ln in st.lines)
    return st


def _parse_csv(data: bytes) -> ParsedStatement:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";" if sample.count(";") >= sample.count(",") else ","
    rows = [r for r in csv.reader(io.StringIO(text), delimiter=delimiter)]
    return _rows_to_statement(rows, "csv")


def _parse_xlsx(data: bytes) -> ParsedStatement:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - archivo corrupto o no es Excel
        raise BusinessRuleViolation("No se pudo leer el archivo Excel") from exc
    ws = wb.worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True, max_row=20000)]
    return _rows_to_statement(rows, "xlsx")


def _ofx_tag(block: str, tag: str) -> str | None:
    m = re.search(rf"<{tag}>([^<\r\n]*)", block, re.IGNORECASE)
    return m.group(1).strip() if m else None


def _parse_ofx(data: bytes) -> ParsedStatement:
    text = data.decode("latin-1")
    st = ParsedStatement(format="ofx")
    for block in re.findall(r"<STMTTRN>(.*?)</STMTTRN>", text, re.IGNORECASE | re.DOTALL):
        d = parse_date((_ofx_tag(block, "DTPOSTED") or "")[:8])
        amount = parse_number(_ofx_tag(block, "TRNAMT"))
        if d is None or not amount:
            continue
        desc = _ofx_tag(block, "NAME") or _ofx_tag(block, "MEMO") or "Movimiento bancario"
        ref = _ofx_tag(block, "CHECKNUM") or _ofx_tag(block, "FITID")
        st.lines.append({"txn_date": d, "description": desc[:255], "reference": (ref or "")[:80] or None,
                         "amount": amount})
    if not st.lines:
        raise BusinessRuleViolation("El archivo OFX no contiene movimientos")
    ledger = re.search(r"<LEDGERBAL>(.*?)</LEDGERBAL>", text, re.IGNORECASE | re.DOTALL)
    if ledger:
        st.closing_balance = parse_number(_ofx_tag(ledger.group(1), "BALAMT"))
        st.statement_date = parse_date((_ofx_tag(ledger.group(1), "DTASOF") or "")[:8])
    st.statement_date = st.statement_date or max(ln["txn_date"] for ln in st.lines)
    if st.closing_balance is not None:
        st.opening_balance = st.closing_balance - sum((ln["amount"] for ln in st.lines), Decimal("0"))
    st.lines.sort(key=lambda ln: ln["txn_date"])
    return st


def parse_statement(filename: str, data: bytes) -> dict:
    if len(data) > MAX_BYTES:
        raise BusinessRuleViolation("El archivo supera 5 MB")
    name = (filename or "").lower()
    head = data[:200].lstrip().upper()
    if name.endswith((".ofx", ".qfx")) or head.startswith(b"OFXHEADER") or b"<OFX>" in data[:2000].upper():
        st = _parse_ofx(data)
    elif name.endswith((".xlsx", ".xlsm")) or data[:2] == b"PK":
        st = _parse_xlsx(data)
    elif name.endswith(".xls"):
        raise BusinessRuleViolation("Formato Excel antiguo (.xls): guárdelo como .xlsx o CSV")
    else:
        st = _parse_csv(data)
    if st.opening_balance is not None and st.closing_balance is not None:
        net = sum((ln["amount"] for ln in st.lines), Decimal("0"))
        if st.opening_balance + net != st.closing_balance:
            st.warnings.append("Saldo inicial + movimientos no coincide con el saldo final del archivo; revise los saldos")
    return st.as_dict()
