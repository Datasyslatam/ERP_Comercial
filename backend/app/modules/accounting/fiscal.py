"""Libros y reportes fiscales.

Venezuela: libro de ventas y libro de compras de IVA, comprobante de retención de IVA y archivo TXT de
retenciones para el portal del SENIAT (Providencia Administrativa SNAT/2015/0049).
Colombia: auxiliar de IVA (generado, descontable y retenciones) y certificados de retención en la fuente,
de IVA y de ICA por tercero y periodo.

Todo se expresa en la moneda funcional de la empresa (bolívares o pesos), convertido con la tasa del documento.
Los formatos deben validarse con el contador antes de presentarse: el ERP los genera, no los declara.
"""
import io
import re
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.errors import BusinessRuleViolation, NotFound
from app.core.money import ZERO, money, to_functional
from app.modules.accounting.models import OpenItem, Payment, PaymentAllocation, PaymentWithholding
from app.modules.einvoicing.models import SalesInvoice
from app.modules.masters.models import Party, Tax
from app.modules.org.models import Company
from app.modules.purchasing.models import SupplierInvoice, SupplierInvoiceTax

WITHHOLDING_KINDS = {"income": "withholding_income", "vat": "withholding_vat", "ica": "withholding_ica"}
DOC_TYPES = {"invoice": "Factura", "credit_note": "Nota crédito", "debit_note": "Nota débito"}


# ---------------------------------------------------------------- utilidades
def rif(id_type: str | None, tax_id: str | None, check_digit: str | None = None) -> str:
    """RIF sin guiones ni espacios (J123456789), como lo exige el TXT del SENIAT."""
    s = re.sub(r"[^A-Za-z0-9]", "", tax_id or "").upper()
    if s and s[0].isdigit():
        s = "V" + s
    if check_digit and len(s) < 10:
        s += check_digit
    return s


def company_rif(c: Company) -> str:
    return rif("RIF", c.tax_id, c.tax_id_check_digit)


def party_rif(party: Party) -> str:
    return rif(party.id_type, party.tax_id, party.check_digit)


def _require_country(company: Company, country: str) -> None:
    if company.country_code != country:
        name = {"VE": "Venezuela", "CO": "Colombia"}[country]
        raise BusinessRuleViolation(f"Este reporte aplica a empresas de {name}")


def _check_range(start: date, end: date) -> None:
    if end < start:
        raise BusinessRuleViolation("La fecha final no puede ser anterior a la inicial")
    if (end - start).days > 400:
        raise BusinessRuleViolation("El rango máximo de un reporte fiscal es de un año")


def _parties(db: Session, ids) -> dict[int, Party]:
    ids = {i for i in ids if i}
    return {p.id: p for p in db.scalars(select(Party).where(Party.id.in_(ids)))} if ids else {}


def _taxes(db: Session, company_id: int) -> dict[int, Tax]:
    return {t.id: t for t in db.scalars(select(Tax).where(Tax.company_id == company_id))}


def _bucket(rate: Decimal) -> str:
    """Alícuota del IVA venezolano: general (16 %), reducida (8 %) o general + adicional (bienes suntuarios)."""
    if rate < 16:
        return "reduced"
    return "general" if rate == 16 else "additional"


def out(value):
    """Decimales como texto (sin pérdida de precisión) para la respuesta JSON."""
    if isinstance(value, Decimal):
        return str(money(value))
    if isinstance(value, dict):
        return {k: out(v) for k, v in value.items()}
    if isinstance(value, list):
        return [out(v) for v in value]
    if isinstance(value, date):
        return value.isoformat()
    return value


@dataclass
class VatSplit:
    """Desglose del IVA de un documento en moneda funcional."""

    total: Decimal = ZERO
    exempt: Decimal = ZERO
    general_base: Decimal = ZERO
    general_tax: Decimal = ZERO
    reduced_base: Decimal = ZERO
    reduced_tax: Decimal = ZERO
    additional_base: Decimal = ZERO
    additional_tax: Decimal = ZERO
    by_rate: dict = field(default_factory=dict)  # tarifa → [base, impuesto]

    @property
    def vat(self) -> Decimal:
        return self.general_tax + self.reduced_tax + self.additional_tax


def _split(doc, taxes: dict[int, Tax], sign: int = 1) -> VatSplit:
    fx = doc.exchange_rate
    s = VatSplit(total=sign * to_functional(doc.total, fx))
    taxed = ZERO
    for t in doc.taxes:
        tax = taxes[t.tax_id]
        if getattr(t, "is_withholding", False) or tax.kind != "vat" or t.rate <= 0:
            continue
        base, amount = sign * to_functional(t.base, fx), sign * to_functional(t.amount, fx)
        b = _bucket(t.rate)
        setattr(s, f"{b}_base", getattr(s, f"{b}_base") + base)
        setattr(s, f"{b}_tax", getattr(s, f"{b}_tax") + amount)
        acc = s.by_rate.setdefault(t.rate, [ZERO, ZERO])
        acc[0] += base
        acc[1] += amount
        taxed += base
    # Exento / no sujeto: todo lo que no es base gravada ni impuesto
    s.exempt = s.total - sign * to_functional(doc.tax_total, fx) - taxed
    return s


_SPLIT_FIELDS = ("total", "exempt", "general_base", "general_tax", "reduced_base", "reduced_tax", "additional_base",
                 "additional_tax")


def _totals(rows: list[dict], extra: tuple[str, ...] = ("withheld",)) -> dict:
    return {f: sum((r.get(f) or ZERO for r in rows), ZERO) for f in _SPLIT_FIELDS + extra}


# ---------------------------------------------------------------- retenciones de IVA practicadas (Venezuela)
@dataclass
class VatWithholding:
    certificate: str | None
    on: date
    invoice: SupplierInvoice
    tax: Tax
    base: Decimal  # IVA sobre el que se retiene, en moneda funcional
    amount: Decimal  # retenido, en moneda funcional
    status: str  # posted / voided
    source: str  # supplier_invoice / payment
    source_id: int


def our_vat_withholdings(db: Session, company_id: int, start: date | None = None, end: date | None = None,
                         number: str | None = None) -> list[VatWithholding]:
    """Retenciones de IVA que la empresa practicó a proveedores, al causar la factura o al pagarla."""
    out_: list[VatWithholding] = []
    q1 = (select(SupplierInvoiceTax, SupplierInvoice, Tax).join(SupplierInvoice).join(Tax, Tax.id == SupplierInvoiceTax.tax_id)
          .where(SupplierInvoice.company_id == company_id, SupplierInvoice.status == "posted",
                 SupplierInvoiceTax.is_withholding.is_(True), Tax.kind == "withholding_vat"))
    q2 = (select(PaymentWithholding, Payment, Tax).join(Payment).join(Tax, Tax.id == PaymentWithholding.tax_id)
          .where(Payment.company_id == company_id, Payment.kind == "disbursement", Tax.kind == "withholding_vat"))
    if start:
        q1 = q1.where(SupplierInvoiceTax.certificate_date >= start)
        q2 = q2.where(Payment.payment_date >= start)
    if end:
        q1 = q1.where(SupplierInvoiceTax.certificate_date <= end)
        q2 = q2.where(Payment.payment_date <= end)
    if number:
        q1 = q1.where(SupplierInvoiceTax.certificate_number == number)
        q2 = q2.where(PaymentWithholding.certificate_number == number)
    for t, inv, tax in db.execute(q1).all():
        out_.append(VatWithholding(t.certificate_number, t.certificate_date or inv.invoice_date, inv, tax,
                                   to_functional(t.base, inv.exchange_rate), to_functional(t.amount, inv.exchange_rate),
                                   "posted", "supplier_invoice", inv.id))
    for w, pay, tax in db.execute(q2).all():
        alloc = db.scalar(select(OpenItem).join(PaymentAllocation, PaymentAllocation.open_item_id == OpenItem.id)
                          .where(PaymentAllocation.payment_id == pay.id, OpenItem.document_type == "supplier_invoice")
                          .order_by(PaymentAllocation.id).limit(1))
        inv = db.get(SupplierInvoice, alloc.document_id) if alloc else None
        if inv is None:
            continue
        out_.append(VatWithholding(w.certificate_number, pay.payment_date, inv, tax, to_functional(w.base, pay.exchange_rate),
                                   to_functional(w.amount, pay.exchange_rate), pay.status, "payment", pay.id))
    out_.sort(key=lambda w: (w.on, w.certificate or ""))
    return out_


# ---------------------------------------------------------------- libros de IVA (Venezuela)
def _row_from_split(s: VatSplit) -> dict:
    return {f: getattr(s, f) for f in _SPLIT_FIELDS}


def ve_sales_book(db: Session, company: Company, start: date, end: date) -> dict:
    _require_country(company, "VE")
    _check_range(start, end)
    taxes = _taxes(db, company.id)
    invoices = list(db.scalars(select(SalesInvoice).where(
        SalesInvoice.company_id == company.id, SalesInvoice.status.in_(("issued", "cancelled")),
        SalesInvoice.issue_date >= start, SalesInvoice.issue_date <= end)
        .order_by(SalesInvoice.issue_date, SalesInvoice.id)))
    refs = {i.id: i for i in db.scalars(select(SalesInvoice).where(
        SalesInvoice.id.in_({i.reference_invoice_id for i in invoices if i.reference_invoice_id})))} \
        if any(i.reference_invoice_id for i in invoices) else {}
    # Retenciones de IVA que nos practicaron los clientes (comprobante recibido en el periodo)
    ret = db.execute(select(PaymentWithholding, Payment).join(Payment).join(Tax, Tax.id == PaymentWithholding.tax_id)
                     .where(Payment.company_id == company.id, Payment.kind == "receipt", Payment.status == "posted",
                            Tax.kind == "withholding_vat", Payment.payment_date >= start, Payment.payment_date <= end)
                     .order_by(Payment.payment_date, Payment.id)).all()
    parties = _parties(db, [i.customer_id for i in invoices] + [p.party_id for _, p in ret])
    rows, warnings = [], []
    for inv in invoices:
        party = parties[inv.customer_id]
        cancelled = inv.status == "cancelled"
        sign = -1 if inv.doc_type == "credit_note" else 1
        row = {"date": inv.issue_date, "kind": inv.doc_type, "rif": party_rif(party), "name": party.legal_name,
               "taxpayer": "Contribuyente" if party.id_type == "RIF" else "No contribuyente",
               "invoice_number": inv.full_number if inv.doc_type == "invoice" else None,
               "control_number": inv.control_number,
               "debit_note": inv.full_number if inv.doc_type == "debit_note" else None,
               "credit_note": inv.full_number if inv.doc_type == "credit_note" else None,
               "transaction": "03-ANU" if cancelled else "01-REG",
               "affected": refs[inv.reference_invoice_id].full_number if inv.reference_invoice_id in refs else None,
               "withheld": ZERO, "certificate": None, "document_type": "sales_invoice", "document_id": inv.id}
        row.update({f: ZERO for f in _SPLIT_FIELDS} if cancelled else _row_from_split(_split(inv, taxes, sign)))
        if not inv.control_number:
            warnings.append(f"{inv.full_number}: sin número de control de la imprenta digital")
        rows.append(row)
    for w, pay in ret:
        party = parties[pay.party_id]
        item = db.scalar(select(OpenItem).join(PaymentAllocation, PaymentAllocation.open_item_id == OpenItem.id)
                         .where(PaymentAllocation.payment_id == pay.id).order_by(PaymentAllocation.id).limit(1))
        rows.append({"date": pay.payment_date, "kind": "withholding", "rif": party_rif(party), "name": party.legal_name,
                     "taxpayer": "Contribuyente", "invoice_number": None, "control_number": None, "debit_note": None,
                     "credit_note": None, "transaction": "01-REG", "affected": item.document_number if item else None,
                     **{f: ZERO for f in _SPLIT_FIELDS}, "withheld": to_functional(w.amount, pay.exchange_rate),
                     "certificate": w.certificate_number, "document_type": "payment", "document_id": pay.id})
        if not w.certificate_number:
            warnings.append(f"Recaudo {pay.number}: retención de IVA sin número de comprobante del cliente")
    rows.sort(key=lambda r: (r["date"], r["kind"] == "withholding"))
    for n, r in enumerate(rows, 1):
        r["n"] = n
    return {"company": company.legal_name, "rif": company_rif(company), "start": start, "end": end,
            "currency": company.functional_currency, "rows": rows, "totals": _totals(rows), "warnings": warnings}


def ve_purchase_book(db: Session, company: Company, start: date, end: date) -> dict:
    _require_country(company, "VE")
    _check_range(start, end)
    taxes = _taxes(db, company.id)
    invoices = list(db.scalars(select(SupplierInvoice).where(
        SupplierInvoice.company_id == company.id, SupplierInvoice.status == "posted",
        SupplierInvoice.invoice_date >= start, SupplierInvoice.invoice_date <= end)
        .order_by(SupplierInvoice.invoice_date, SupplierInvoice.id)))
    withholdings = [w for w in our_vat_withholdings(db, company.id, start, end) if w.status == "posted"]
    by_invoice = defaultdict(list)
    for w in withholdings:
        by_invoice[w.invoice.id].append(w)
    in_period = {i.id for i in invoices}
    parties = _parties(db, [i.supplier_id for i in invoices] + [w.invoice.supplier_id for w in withholdings])
    rows, warnings = [], []

    def base_row(inv: SupplierInvoice) -> dict:
        party = parties[inv.supplier_id]
        return {"rif": party_rif(party), "name": party.legal_name, "invoice_number": inv.supplier_invoice_number,
                "control_number": inv.supplier_control_number, "debit_note": None, "credit_note": None,
                "transaction": "01-REG", "affected": None, "document_type": "supplier_invoice", "document_id": inv.id}

    for inv in invoices:
        ws = by_invoice.get(inv.id, [])
        row = {"date": inv.invoice_date, "kind": "invoice", **base_row(inv), **_row_from_split(_split(inv, taxes)),
               "withheld": sum((w.amount for w in ws), ZERO),
               "certificate": ", ".join(w.certificate for w in ws if w.certificate) or None,
               "certificate_date": ws[0].on if ws else None}
        if not inv.supplier_control_number:
            warnings.append(f"Factura {inv.supplier_invoice_number} ({inv.number}): sin número de control")
        rows.append(row)
    # Retenciones del periodo sobre facturas registradas en periodos anteriores
    for w in withholdings:
        if w.invoice.id in in_period:
            continue
        rows.append({"date": w.on, "kind": "withholding", **base_row(w.invoice), "affected": w.invoice.supplier_invoice_number,
                     **{f: ZERO for f in _SPLIT_FIELDS}, "withheld": w.amount, "certificate": w.certificate,
                     "certificate_date": w.on})
    rows.sort(key=lambda r: (r["date"], r["kind"] == "withholding"))
    for n, r in enumerate(rows, 1):
        r["n"] = n
    return {"company": company.legal_name, "rif": company_rif(company), "start": start, "end": end,
            "currency": company.functional_currency, "rows": rows, "totals": _totals(rows), "warnings": warnings}


# ---------------------------------------------------------------- comprobante de retención y TXT (Venezuela)
def ve_vat_certificates(db: Session, company: Company, start: date, end: date) -> list[dict]:
    _require_country(company, "VE")
    _check_range(start, end)
    ws = our_vat_withholdings(db, company.id, start, end)
    parties = _parties(db, [w.invoice.supplier_id for w in ws])
    return [{"certificate": w.certificate, "date": w.on, "supplier": parties[w.invoice.supplier_id].legal_name,
             "rif": party_rif(parties[w.invoice.supplier_id]), "invoice_number": w.invoice.supplier_invoice_number,
             "control_number": w.invoice.supplier_control_number, "vat": w.base, "rate": w.tax.rate,
             "withheld": w.amount, "status": w.status, "source": w.source, "source_id": w.source_id} for w in ws]


def _certificate_lines(db: Session, company: Company, w: VatWithholding, taxes: dict[int, Tax]) -> list[dict]:
    """Una línea por alícuota de la factura; lo retenido se reparte en proporción al IVA de cada alícuota."""
    s = _split(w.invoice, taxes)
    vat = s.vat or Decimal("1")
    lines, assigned = [], ZERO
    rates = sorted(s.by_rate.items())
    for i, (rate, (base, tax)) in enumerate(rates):
        withheld = w.amount - assigned if i == len(rates) - 1 else money(w.amount * tax / vat)
        assigned += withheld
        lines.append({"rate": rate, "base": base, "tax": tax, "withheld": withheld,
                      "total": s.total if i == 0 else ZERO, "exempt": s.exempt if i == 0 else ZERO})
    return lines or [{"rate": ZERO, "base": ZERO, "tax": ZERO, "withheld": w.amount, "total": s.total,
                      "exempt": s.exempt}]


def ve_vat_certificate_pdf(db: Session, company: Company, number: str) -> tuple[str, bytes]:
    from app.modules.documents.pdf import DocData, fdate, render
    from app.modules.documents.pdf import money as fmt
    from app.modules.documents.service import _company_block, _party_block

    _require_country(company, "VE")
    ws = our_vat_withholdings(db, company.id, number=number)
    if not ws:
        raise NotFound("Comprobante de retención no encontrado")
    w = ws[0]
    supplier = db.get(Party, w.invoice.supplier_id)
    taxes = _taxes(db, company.id)
    rows, total_withheld = [], ZERO
    for n, ln in enumerate(_certificate_lines(db, company, w, taxes), 1):
        rows.append([str(n), f"Factura {w.invoice.supplier_invoice_number} · Control {w.invoice.supplier_control_number or '-'}"
                     f" · 01-REG", fdate(w.invoice.invoice_date), fmt(ln["total"]) if n == 1 else "-",
                     fmt(ln["exempt"]) if n == 1 else "-", fmt(ln["base"]),
                     f"{fmt(ln['rate'])} %", fmt(ln["tax"]), fmt(ln["withheld"])])
        total_withheld += ln["withheld"]
    data = DocData(
        title="Comprobante de retención de IVA", number=number, company=_company_block(company),
        party_label="Sujeto retenido", party=_party_block(supplier),
        fields=[("Fecha de emisión", fdate(w.on)), ("Período fiscal", f"Año {w.on:%Y} · Mes {w.on:%m}"),
                ("RIF del agente", company_rif(company)), ("RIF del retenido", party_rif(supplier)),
                ("Retención", f"{fmt(w.tax.rate)} % del IVA"), ("Moneda", company.functional_currency)],
        lines=[], columns=[("Op.", 8, "C"), ("Documento", 36, "L"), ("Fecha", 18, "C"), ("Total c/IVA", 24, "R"),
                           ("Exento", 18, "R"), ("Base imp.", 24, "R"), ("Alíc.", 15, "R"), ("IVA", 20, "R"),
                           ("Retenido", 24.9, "R")],
        rows=rows, totals=[("IVA retenido", fmt(total_withheld), True)], currency=company.functional_currency,
        notes="Comprobante emitido conforme a la Providencia Administrativa SNAT/2015/0049 sobre agentes de "
              "retención del IVA.",
        notice="ANULADO" if w.status == "voided" else None,
        terms="Agente de retención: ______________________________\n"
              "Recibido por (sujeto retenido): ____________________\nFecha de recepción: ____________")
    return f"Comprobante_retencion_IVA_{number}.pdf", render(data)


def ve_vat_txt(db: Session, company: Company, year: int, month: int, half: int) -> tuple[str, str, list[str]]:
    """Archivo TXT de retenciones de IVA (una línea por documento y alícuota, separadas por tabulador).

    Campos: RIF del agente, período (AAAAMM), fecha del documento, tipo de operación (C), tipo de documento
    (01 factura), RIF del proveedor, número del documento, número de control, total del documento, base
    imponible, IVA retenido, documento afectado, número de comprobante, monto exento, alícuota y expediente.
    half: 1 = días 1 a 15, 2 = del 16 al fin de mes, 0 = mes completo.
    """
    _require_country(company, "VE")
    if not 1 <= month <= 12 or half not in (0, 1, 2):
        raise BusinessRuleViolation("Período inválido")
    import calendar

    last = calendar.monthrange(year, month)[1]
    start = date(year, month, 16 if half == 2 else 1)
    end = date(year, month, 15 if half == 1 else last)
    taxes = _taxes(db, company.id)
    agent = company_rif(company)
    lines, warnings = [], []
    for w in our_vat_withholdings(db, company.id, start, end):
        if w.status != "posted":
            continue
        inv = w.invoice
        supplier = db.get(Party, inv.supplier_id)
        if not w.certificate or not inv.supplier_control_number:
            warnings.append(f"Factura {inv.supplier_invoice_number}: falta número de control o de comprobante")
        cert_lines = _certificate_lines(db, company, w, taxes)
        doc_total = cert_lines[0]["total"]  # el total es del documento: se repite en cada alícuota; el exento, no
        for ln in cert_lines:
            lines.append("\t".join([
                agent, f"{w.on:%Y%m}", inv.invoice_date.isoformat(), "C", "01", party_rif(supplier),
                inv.supplier_invoice_number, inv.supplier_control_number or "", f"{doc_total:.2f}", f"{ln['base']:.2f}",
                f"{ln['withheld']:.2f}", "0", w.certificate or "", f"{ln['exempt']:.2f}", f"{ln['rate']:.2f}", "0"]))
    name = f"retenciones_iva_{agent}_{year}{month:02d}{'' if half == 0 else f'_q{half}'}.txt"
    return name, "\r\n".join(lines) + ("\r\n" if lines else ""), warnings


# ---------------------------------------------------------------- Colombia: auxiliar de IVA
def co_vat_auxiliary(db: Session, company: Company, start: date, end: date) -> dict:
    _require_country(company, "CO")
    _check_range(start, end)
    taxes = _taxes(db, company.id)
    sales = list(db.scalars(select(SalesInvoice).where(
        SalesInvoice.company_id == company.id, SalesInvoice.status == "issued",
        SalesInvoice.issue_date >= start, SalesInvoice.issue_date <= end).order_by(SalesInvoice.issue_date, SalesInvoice.id)))
    purchases = list(db.scalars(select(SupplierInvoice).where(
        SupplierInvoice.company_id == company.id, SupplierInvoice.status == "posted",
        SupplierInvoice.invoice_date >= start, SupplierInvoice.invoice_date <= end)
        .order_by(SupplierInvoice.invoice_date, SupplierInvoice.id)))
    received = db.execute(select(PaymentWithholding, Payment).join(Payment).join(Tax, Tax.id == PaymentWithholding.tax_id)
                          .where(Payment.company_id == company.id, Payment.kind == "receipt", Payment.status == "posted",
                                 Tax.kind == "withholding_vat", Payment.payment_date >= start, Payment.payment_date <= end)
                          .order_by(Payment.payment_date)).all()
    parties = _parties(db, [i.customer_id for i in sales] + [i.supplier_id for i in purchases]
                       + [p.party_id for _, p in received])

    def vat_rows(doc, party_id, number, label, on, sign):
        party = parties[party_id]
        for t in doc.taxes:
            tax = taxes[t.tax_id]
            if getattr(t, "is_withholding", False) or tax.kind != "vat":
                continue
            yield {"date": on, "document": label, "number": number, "nit": party.tax_id, "name": party.legal_name,
                   "rate": t.rate, "base": sign * to_functional(t.base, doc.exchange_rate),
                   "tax": sign * to_functional(t.amount, doc.exchange_rate)}

    generated = [r for inv in sales for r in vat_rows(inv, inv.customer_id, inv.full_number, DOC_TYPES[inv.doc_type],
                                                      inv.issue_date, -1 if inv.doc_type == "credit_note" else 1)]
    deductible = [r for inv in purchases for r in vat_rows(inv, inv.supplier_id, inv.supplier_invoice_number,
                                                           "Factura de proveedor", inv.invoice_date, 1)]
    withheld_from_us = [{"date": pay.payment_date, "document": "Recaudo", "number": pay.number,
                         "nit": parties[pay.party_id].tax_id, "name": parties[pay.party_id].legal_name,
                         "certificate": w.certificate_number, "tax": to_functional(w.amount, pay.exchange_rate)}
                        for w, pay in received]

    def by_rate(rows):
        acc = defaultdict(lambda: [ZERO, ZERO])
        for r in rows:
            acc[r["rate"]][0] += r["base"]
            acc[r["rate"]][1] += r["tax"]
        return [{"rate": k, "base": v[0], "tax": v[1]} for k, v in sorted(acc.items(), reverse=True)]

    g, d = sum((r["tax"] for r in generated), ZERO), sum((r["tax"] for r in deductible), ZERO)
    w = sum((r["tax"] for r in withheld_from_us), ZERO)
    return {"company": company.legal_name, "nit": f"{company.tax_id}-{company.tax_id_check_digit or ''}".rstrip("-"),
            "start": start, "end": end, "currency": company.functional_currency,
            "generated": {"rows": generated, "by_rate": by_rate(generated), "total": g},
            "deductible": {"rows": deductible, "by_rate": by_rate(deductible), "total": d},
            "withheld_from_us": {"rows": withheld_from_us, "total": w},
            "summary": {"generated": g, "deductible": d, "balance": g - d, "withheld_from_us": w,
                        "payable": g - d - w}}


# ---------------------------------------------------------------- Colombia: certificados de retención
CERT_TITLES = {"income": "Certificado de retención en la fuente", "vat": "Certificado de retención de IVA",
               "ica": "Certificado de retención de ICA"}
CERT_PREFIXES = {"income": "RF", "vat": "RIVA", "ica": "RICA"}


def co_withholdings(db: Session, company: Company, start: date, end: date, kind: str,
                    party_id: int | None = None) -> list[dict]:
    """Retenciones practicadas a proveedores en el periodo, agrupadas por tercero y concepto."""
    _require_country(company, "CO")
    _check_range(start, end)
    if kind not in WITHHOLDING_KINDS:
        raise BusinessRuleViolation("Tipo de retención inválido (income, vat o ica)")
    tax_kind = WITHHOLDING_KINDS[kind]
    q1 = (select(SupplierInvoiceTax, SupplierInvoice, Tax).join(SupplierInvoice).join(Tax, Tax.id == SupplierInvoiceTax.tax_id)
          .where(SupplierInvoice.company_id == company.id, SupplierInvoice.status == "posted",
                 SupplierInvoiceTax.is_withholding.is_(True), Tax.kind == tax_kind,
                 SupplierInvoice.invoice_date >= start, SupplierInvoice.invoice_date <= end))
    q2 = (select(PaymentWithholding, Payment, Tax).join(Payment).join(Tax, Tax.id == PaymentWithholding.tax_id)
          .where(Payment.company_id == company.id, Payment.kind == "disbursement", Payment.status == "posted",
                 Tax.kind == tax_kind, Payment.payment_date >= start, Payment.payment_date <= end))
    if party_id:
        q1 = q1.where(SupplierInvoice.supplier_id == party_id)
        q2 = q2.where(Payment.party_id == party_id)
    acc: dict[int, dict[int, list]] = defaultdict(lambda: defaultdict(lambda: [ZERO, ZERO]))
    for t, inv, tax in db.execute(q1).all():
        a = acc[inv.supplier_id][tax.id]
        a[0] += to_functional(t.base, inv.exchange_rate)
        a[1] += to_functional(t.amount, inv.exchange_rate)
    for w, pay, tax in db.execute(q2).all():
        a = acc[pay.party_id][tax.id]
        a[0] += to_functional(w.base, pay.exchange_rate)
        a[1] += to_functional(w.amount, pay.exchange_rate)
    taxes = _taxes(db, company.id)
    parties = _parties(db, acc.keys())
    result = []
    for pid, by_tax in sorted(acc.items(), key=lambda kv: parties[kv[0]].legal_name):
        lines = [{"code": taxes[tid].code, "concept": taxes[tid].name, "rate": taxes[tid].rate, "base": v[0], "amount": v[1]}
                 for tid, v in sorted(by_tax.items())]
        party = parties[pid]
        result.append({"party_id": pid, "nit": party.tax_id, "check_digit": party.check_digit, "name": party.legal_name,
                       "email": party.email, "lines": lines, "base": sum((ln["base"] for ln in lines), ZERO),
                       "amount": sum((ln["amount"] for ln in lines), ZERO)})
    return result


def co_certificate_pdf(db: Session, company: Company, party_id: int, start: date, end: date,
                       kind: str) -> tuple[str, bytes]:
    from app.modules.documents.pdf import DocData, fdate, render
    from app.modules.documents.pdf import money as fmt
    from app.modules.documents.service import _company_block, _party_block
    from app.modules.documents.words import amount_in_words

    data = co_withholdings(db, company, start, end, kind, party_id)
    if not data:
        raise NotFound("El tercero no tiene retenciones de ese tipo en el periodo")
    cert = data[0]
    party = db.get(Party, party_id)
    city = company.city
    if kind == "ica":
        where = f"al municipio de {city}" if city else "al municipio correspondiente"
    else:
        where = f"a la DIAN en la ciudad de {city}" if city else "a la DIAN"
    number = f"{CERT_PREFIXES[kind]}-{start:%Y%m%d}-{end:%Y%m%d}-{party.tax_id}"
    doc = DocData(
        title=CERT_TITLES[kind], number=number, company=_company_block(company), party_label="Retenido",
        party=_party_block(party),
        fields=[("Periodo", f"{fdate(start)} a {fdate(end)}"), ("Expedido el", fdate(date.today())),
                ("Ciudad de práctica", city or "-"), ("Ciudad de consignación", city or "-")],
        lines=[], columns=[("Código", 22, "L"), ("Concepto", 75.9, "L"), ("Tarifa", 18, "R"),
                           ("Base" if kind != "vat" else "Base (IVA)", 36, "R"), ("Retenido", 36, "R")],
        rows=[[ln["code"], ln["concept"], f"{fmt(ln['rate'], 3)} %", fmt(ln["base"]), fmt(ln["amount"])]
              for ln in cert["lines"]],
        totals=[("Total base", fmt(cert["base"]), False), ("Total retenido", fmt(cert["amount"]), True)],
        currency=company.functional_currency,
        notes=f"Valor retenido: {amount_in_words(cert['amount'], company.functional_currency)}",
        terms=f"Certificado expedido conforme al artículo 381 del Estatuto Tributario. Los valores retenidos fueron "
              f"declarados y consignados oportunamente {where}. No requiere firma autógrafa (artículo 10 del "
              f"Decreto 836 de 1991).")
    slug = re.sub(r"[^A-Za-z0-9]+", "_", CERT_TITLES[kind]).strip("_")
    return f"{slug}_{party.tax_id}_{start:%Y%m%d}_{end:%Y%m%d}.pdf", render(doc)


# ---------------------------------------------------------------- exportación a Excel
VE_SALES_COLUMNS = [
    ("n", "N.º"), ("date", "Fecha"), ("rif", "RIF"), ("name", "Nombre o razón social"), ("taxpayer", "Tipo de cliente"),
    ("invoice_number", "N.º factura"), ("control_number", "N.º control"), ("debit_note", "N.º nota débito"),
    ("credit_note", "N.º nota crédito"), ("transaction", "Tipo de transacción"), ("affected", "Factura afectada"),
    ("total", "Total ventas con IVA"), ("exempt", "Ventas exentas o no sujetas"),
    ("general_base", "Base alícuota general"), ("general_tax", "IVA alícuota general"),
    ("reduced_base", "Base alícuota reducida"), ("reduced_tax", "IVA alícuota reducida"),
    ("additional_base", "Base alícuota general + adicional"), ("additional_tax", "IVA alícuota general + adicional"),
    ("withheld", "IVA retenido por el comprador"), ("certificate", "N.º comprobante de retención"),
]
VE_PURCHASE_COLUMNS = [
    ("n", "N.º"), ("date", "Fecha"), ("rif", "RIF"), ("name", "Nombre o razón social del proveedor"),
    ("invoice_number", "N.º factura"), ("control_number", "N.º control"), ("debit_note", "N.º nota débito"),
    ("credit_note", "N.º nota crédito"), ("transaction", "Tipo de transacción"), ("affected", "Factura afectada"),
    ("total", "Total compras con IVA"), ("exempt", "Compras exentas o sin derecho a crédito"),
    ("general_base", "Base alícuota general"), ("general_tax", "IVA alícuota general"),
    ("reduced_base", "Base alícuota reducida"), ("reduced_tax", "IVA alícuota reducida"),
    ("additional_base", "Base alícuota general + adicional"), ("additional_tax", "IVA alícuota general + adicional"),
    ("withheld", "IVA retenido al vendedor"), ("certificate", "N.º comprobante de retención"),
    ("certificate_date", "Fecha del comprobante"),
]
CO_AUX_COLUMNS = [("date", "Fecha"), ("document", "Documento"), ("number", "Número"), ("nit", "NIT"),
                  ("name", "Tercero"), ("rate", "Tarifa %"), ("base", "Base"), ("tax", "IVA")]


def xlsx(title: str, company: str, tax_id: str, period: str, sections: list[tuple[str | None, list, list[dict], dict | None]]
         ) -> bytes:
    """sections: (subtítulo, columnas [(clave, título)], filas, totales)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    bold, head_fill = Font(bold=True), PatternFill("solid", fgColor="0A192F")
    ws.append([company])
    ws.append([tax_id])
    ws.append([title])
    ws.append([period])
    for cell in ("A1", "A3"):
        ws[cell].font = Font(bold=True, size=12)
    widths: dict[int, int] = {}
    for subtitle, columns, rows, totals in sections:
        ws.append([])
        if subtitle:
            ws.append([subtitle])
            ws.cell(ws.max_row, 1).font = bold
        ws.append([c[1] for c in columns])
        for i in range(1, len(columns) + 1):
            c = ws.cell(ws.max_row, i)
            c.font, c.fill = Font(bold=True, color="FFFFFF"), head_fill
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for r in rows:
            ws.append([float(money(r.get(k))) if isinstance(r.get(k), Decimal) else r.get(k) for k, _ in columns])
        if totals:
            ws.append([("Totales" if i == 0 else (float(money(totals[k])) if k in totals else None))
                       for i, (k, _) in enumerate(columns)])
            for i in range(1, len(columns) + 1):
                ws.cell(ws.max_row, i).font = bold
        for i, (k, label) in enumerate(columns, 1):
            widths[i] = max(widths.get(i, 10), min(40, len(label) + 2), 14 if k in _SPLIT_FIELDS + ("withheld", "tax", "base")
                            else 10)
    for row in ws.iter_rows(min_row=5):
        for c in row:
            if isinstance(c.value, float):
                c.number_format = "#,##0.00"
            elif isinstance(c.value, date):
                c.number_format = "DD/MM/YYYY"
    for i, w in widths.items():
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
